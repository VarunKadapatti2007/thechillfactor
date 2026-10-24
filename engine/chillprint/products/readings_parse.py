"""Read a grower's temperature file in whatever format their thermometer or logger exports.

Excel (.xlsx, .xls), Word tables (.docx), PDF text and any delimited text (.csv, .txt, .tsv, logger
exports) become one table of cells. Built-in rules find the date, low, high or temperature columns
when the headers are clear. When they are not (title rows, odd names, separate date and time
columns, unusual date formats), Claude reads a preview of the table and says which column is which.
Code then applies that layout to every row: Claude never supplies a number.

The result feeds POST /api/orchard (products/orchard.py).
"""
from __future__ import annotations

import csv
import io
import json
import os
import re
from datetime import date, datetime, time

import numpy as np
import pandas as pd

from .point import BadRequest

MAX_BYTES = 8_000_000
MAX_ROWS = 100_000
MODEL = "claude-haiku-5-5"  # naming a few columns is a small job: the cheapest current model is plenty

DATE_RE = re.compile(r"date|day|timestamp|time|fecha", re.I)
TIME_RE = re.compile(r"^\s*(time|hour|hora)\b", re.I)
LOW_RE = re.compile(r"min|low", re.I)
HIGH_RE = re.compile(r"max|high", re.I)
TEMP_RE = re.compile(r"temp|air|°|deg|reading|value|\bt\b", re.I)
NOT_AIR_RE = re.compile(r"dew|soil|leaf|wet|humid|rh\b|batt|volt|wind|rain|solar|pressure", re.I)


# --- 1. Any file -> a table of strings --------------------------------------------------------


def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d %H:%M") if (v.hour or v.minute) else v.strftime("%Y-%m-%d")
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, time):
        return v.strftime("%H:%M")
    if isinstance(v, float):
        return f"{v:.6g}"
    return str(v).strip()


def _numeric_cells(rows: list[list[str]]) -> int:
    return sum(1 for r in rows[:500] for c in r if re.fullmatch(r"-?\d+(\.\d+)?", c or ""))


def _excel(data: bytes) -> tuple[list[list[str]], str]:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    best, best_name = [], ""
    for ws in wb.worksheets:
        rows = [[_cell(v) for v in r] for _, r in zip(range(MAX_ROWS), ws.iter_rows(values_only=True))]
        if _numeric_cells(rows) > _numeric_cells(best):
            best, best_name = rows, ws.title
    return best, best_name


def _xls(data: bytes) -> tuple[list[list[str]], str]:
    import xlrd
    book = xlrd.open_workbook(file_contents=data)
    best, best_name = [], ""
    for sh in book.sheets():
        rows = []
        for i in range(min(sh.nrows, MAX_ROWS)):
            row = []
            for c in sh.row(i):
                if c.ctype == xlrd.XL_CELL_DATE:
                    row.append(_cell(xlrd.xldate_as_datetime(c.value, book.datemode)))
                else:
                    row.append(_cell(c.value))
            rows.append(row)
        if _numeric_cells(rows) > _numeric_cells(best):
            best, best_name = rows, sh.name
    return best, best_name


def _docx(data: bytes) -> list[list[str]]:
    import docx
    d = docx.Document(io.BytesIO(data))
    tables = [[[c.text.strip() for c in r.cells] for r in t.rows] for t in d.tables]
    if tables:
        return max(tables, key=len)
    return _split_lines("\n".join(p.text for p in d.paragraphs))


def _pdf(data: bytes) -> list[list[str]]:
    from pypdf import PdfReader
    text = "\n".join((p.extract_text() or "") for p in PdfReader(io.BytesIO(data)).pages)
    return _split_lines(text)


def _decode(data: bytes) -> str:
    if data[:2] in (b"\xff\xfe", b"\xfe\xff") or data[1:200:2].count(b"\x00") > 50:
        return data.decode("utf-16", errors="replace")
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1", errors="replace")


def _split_lines(text: str) -> list[list[str]]:
    lines = [ln for ln in text.splitlines() if ln.strip()][:MAX_ROWS]
    if not lines:
        return []
    sample = "\n".join(lines[:50])
    try:
        delim = csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
        return [[c.strip() for c in r] for r in csv.reader(lines, delimiter=delim)]
    except csv.Error:
        # Whitespace-aligned exports: split on runs of 2+ spaces, else any whitespace.
        sep = r"\s{2,}" if any(re.search(r"\s{2,}", ln.strip()) for ln in lines[:20]) else r"\s+"
        return [re.split(sep, ln.strip()) for ln in lines]


def load_table(filename: str, data: bytes) -> tuple[list[list[str]], str]:
    """(rows of string cells, a short description of where they came from)."""
    if len(data) > MAX_BYTES:
        raise BadRequest("That file is too large (over 8 MB). Export daily lows and highs, or hourly logger data.")
    ext = os.path.splitext(filename.lower())[1]
    try:
        if ext in (".xlsx", ".xlsm") or data[:2] == b"PK" and ext not in (".docx",):
            rows, sheet = _excel(data)
            return rows, f"Excel sheet '{sheet}'"
        if ext == ".xls":
            rows, sheet = _xls(data)
            return rows, f"Excel sheet '{sheet}'"
        if ext == ".docx":
            return _docx(data), "Word document"
        if ext == ".pdf" or data[:4] == b"%PDF":
            return _pdf(data), "PDF text"
    except BadRequest:
        raise
    except Exception as e:
        raise BadRequest(f"Couldn't open that {ext or 'file'} ({str(e)[:80]}). Try saving it as CSV or Excel.")
    return _split_lines(_decode(data)), "text file"


# --- 2. Which column is which -----------------------------------------------------------------


def _looks_date(s: str) -> bool:
    return bool(re.match(r"^\s*(\d{4}[-/.]\d{1,2}[-/.]\d{1,2}|\d{1,2}[-/.]\d{1,2}[-/.]\d{2,4}|\d{1,2}[ -][A-Za-z]{3})", s or ""))


def _num(s: str) -> float:
    if s is None:
        return np.nan
    v = str(s).strip().replace("°", "").replace("°", "")
    v = re.sub(r"\s*[FfCc]$", "", v)
    if re.fullmatch(r"-?\d+,\d+", v):
        v = v.replace(",", ".")
    try:
        return float(v)
    except ValueError:
        return np.nan


def _unit_from(headers: list[str]) -> str:
    h = " ".join(headers).lower()
    if re.search(r"°\s*c|\(c\)|_c\b|celsius|degc|deg c", h):
        return "C"
    if re.search(r"°\s*f|\(f\)|_f\b|fahrenheit|degf|deg f", h):
        return "F"
    return ""


def _is_index(values: list[str]) -> bool:
    v = np.array([_num(x) for x in values[:50]])
    v = v[np.isfinite(v)]
    return len(v) > 5 and bool(np.all(np.diff(v) == 1))


def rules_mapping(rows: list[list[str]]) -> dict | None:
    """Find the header and columns from their names. None when the layout is not obvious."""
    for hi in range(min(15, len(rows) - 3)):
        head = [c.lower() for c in rows[hi]]
        body = rows[hi + 1:hi + 51]
        if not head or not any(any(_looks_date(c) for c in r[:3]) for r in body[:20]):
            continue
        texty = [c for c in head if c and not re.fullmatch(r"-?[\d.,]+", c) and not _looks_date(c)]
        if len(texty) < 2:
            continue

        def score(i: int, want: re.Pattern) -> int:
            h = head[i]
            if not want.search(h) or NOT_AIR_RE.search(h) or re.search(r"#|\bno\.?\b|index|\bid\b|count", h):
                return -1
            if _is_index(_col(body, i)) or np.isfinite([_num(x) for x in _col(body, i)[:20]]).mean() < 0.5:
                return -1
            return 1 + 3 * ("air" in h) + 2 * ("temp" in h) + ("°" in h or "deg" in h)

        def best(want: re.Pattern, exclude=()) -> int:
            cands = [(score(i, want), i) for i in range(len(head)) if i not in exclude]
            cands = [c for c in cands if c[0] > 0]
            return max(cands)[1] if cands else -1

        d = next((i for i, h in enumerate(head) if DATE_RE.search(h)), -1)
        if d < 0:
            continue
        t = next((i for i, h in enumerate(head) if i != d and TIME_RE.search(h)), -1)
        lo = best(LOW_RE, (d, t))
        hi_ = best(HIGH_RE, (d, t, lo))
        tmp = best(re.compile(TEMP_RE.pattern, re.I), (d, t))
        if tmp >= 0 and (LOW_RE.search(head[tmp]) or HIGH_RE.search(head[tmp])):
            tmp = -1
        if lo >= 0 and hi_ >= 0:
            return {"kind": "daily", "header_rows": hi + 1, "date_column": d, "time_column": -1, "low_column": lo,
                    "high_column": hi_, "temp_column": -1, "date_format": "", "day_first": False,
                    "unit": _unit_from([head[lo], head[hi_]]),
                    "explanation": f"Found a date column and daily low and high columns ('{rows[hi][lo]}', '{rows[hi][hi_]}')."}
        if tmp >= 0:
            return {"kind": "logger", "header_rows": hi + 1, "date_column": d, "time_column": t, "low_column": -1,
                    "high_column": -1, "temp_column": tmp, "date_format": "", "day_first": False,
                    "unit": _unit_from([head[tmp]]),
                    "explanation": f"Found timestamped air temperature readings ('{rows[hi][tmp]}')."}
    return None


SYSTEM = """You identify the layout of temperature records that fruit and nut growers export from \
min/max thermometers, data loggers, weather stations or spreadsheets. You are given a preview of a \
table (row number, then cells separated by ' | '). Say which columns hold what, so code can read \
every row. Do not read or copy any values yourself.

- kind: "daily" when each row is one day with a low (minimum) and a high (maximum) air temperature; \
"logger" when rows are repeated readings with a time and one air temperature; "unusable" when the \
table has no air temperature records.
- Columns are 0-based. Use -1 for a column that does not exist. Use air temperature only, never \
dew point, soil, leaf, humidity or battery columns.
- header_rows: how many rows come before the first data row.
- date_format: a Python strptime pattern for the date column (including the time if it is in the \
same cell), or "" if dates are ISO-like or you are unsure. day_first: true only when dates are \
day/month/year.
- unit: "F" or "C", from the headers or, if they say nothing, from the values (Central Valley \
winter lows and highs are roughly 25-75 F or -4-24 C).
- explanation: one short, plain sentence for the grower saying what you found. Name columns by their \
header text in quotes, never by number."""

SCHEMA = {
    "type": "object",
    "properties": {
        "kind": {"type": "string", "enum": ["daily", "logger", "unusable"]},
        "header_rows": {"type": "integer"},
        "date_column": {"type": "integer"},
        "time_column": {"type": "integer"},
        "low_column": {"type": "integer"},
        "high_column": {"type": "integer"},
        "temp_column": {"type": "integer"},
        "date_format": {"type": "string"},
        "day_first": {"type": "boolean"},
        "unit": {"type": "string", "enum": ["F", "C"]},
        "explanation": {"type": "string"},
    },
    "required": ["kind", "header_rows", "date_column", "time_column", "low_column", "high_column", "temp_column",
                 "date_format", "day_first", "unit", "explanation"],
    "additionalProperties": False,
}


def _api_key() -> str | None:
    """The key is one token starting sk-ant-; anything else (a pasted link, stray text) counts as no key."""
    v = (os.environ.get("ANTHROPIC_API_KEY") or "").split()
    return v[0] if v and v[0].startswith("sk-ant-") else None


def ai_available() -> bool:
    return _api_key() is not None


def _preview(rows: list[list[str]], filename: str, where: str) -> str:
    n = len(rows)
    pick = list(range(min(25, n)))
    if n > 35:
        mid = n // 2
        pick += list(range(mid, mid + 5)) + list(range(n - 5, n))
    out, last = [f"File: {filename} ({where}), {n} rows in total."], -1
    for i in pick:
        if i != last + 1:
            out.append("...")
        out.append(f"{i}: " + " | ".join((c or "")[:40] for c in rows[i][:20]))
        last = i
    return "\n".join(out)


def ai_mapping(rows: list[list[str]], filename: str, where: str) -> dict:
    import anthropic
    client = anthropic.Anthropic(api_key=_api_key(), timeout=60.0, max_retries=2)
    try:
        resp = client.messages.create(
            model=MODEL,
            max_tokens=4000,
            system=SYSTEM,
            output_config={"effort": "low", "format": {"type": "json_schema", "schema": SCHEMA}},
            messages=[{"role": "user", "content": _preview(rows, filename, where)}],
        )
    except anthropic.APIStatusError as e:
        raise BadRequest(f"The file reader is unavailable right now ({e.status_code}). Try a CSV with columns date, low, high.")
    except anthropic.APIConnectionError:
        raise BadRequest("The file reader is unavailable right now. Try a CSV with columns date, low, high.")
    if resp.stop_reason == "refusal":
        raise BadRequest("The file reader could not read this file. Try a CSV with columns date, low, high.")
    text = next((b.text for b in resp.content if b.type == "text"), "")
    return json.loads(text)


# --- 3. Apply a layout to every row -----------------------------------------------------------


def _col(rows: list[list[str]], i: int) -> list[str]:
    return [(r[i] if 0 <= i < len(r) else "") for r in rows]


def _dates(strings: list[str], fmt: str, day_first: bool) -> pd.Series:
    s = pd.Series(strings, dtype="object").fillna("").astype(str).str.strip()
    if fmt:
        out = pd.to_datetime(s, format=fmt, errors="coerce")
        if out.notna().mean() > 0.8:
            return out
    return pd.to_datetime(s, format="mixed", dayfirst=day_first, errors="coerce")


def apply_mapping(rows: list[list[str]], m: dict) -> dict:
    body = rows[max(0, int(m["header_rows"])):]
    if m["kind"] == "unusable" or not body:
        raise BadRequest("No air temperature readings were found in that file.")
    d = _col(body, m["date_column"])
    if m["kind"] == "logger" and m["time_column"] >= 0:
        d = [f"{a} {b}".strip() for a, b in zip(d, _col(body, m["time_column"]))]
    ts = _dates(d, m.get("date_format", ""), bool(m.get("day_first")))
    if m["kind"] == "daily":
        lo = np.array([_num(x) for x in _col(body, m["low_column"])])
        hi = np.array([_num(x) for x in _col(body, m["high_column"])])
        ok = ts.notna().to_numpy() & np.isfinite(lo) & np.isfinite(hi)
        df = pd.DataFrame({"date": ts[ok].dt.strftime("%Y-%m-%d").to_numpy(),
                           "low": np.minimum(lo[ok], hi[ok]), "high": np.maximum(lo[ok], hi[ok])})
        df = df.drop_duplicates("date", keep="last").sort_values("date")
        return {"kind": "daily", "dates": df.date.tolist(), "lows": df.low.round(2).tolist(),
                "highs": df.high.round(2).tolist(), "skipped_rows": int((~ok).sum())}
    t = np.array([_num(x) for x in _col(body, m["temp_column"])])
    ok = ts.notna().to_numpy() & np.isfinite(t)
    has_time = ts[ok].dt.hour.ne(0).any() or ts[ok].dt.minute.ne(0).any()
    if not has_time:
        raise BadRequest("The readings have dates but no times. For a min/max thermometer, use columns: date, low, high.")
    df = pd.DataFrame({"time": ts[ok].dt.strftime("%Y-%m-%d %H:%M").to_numpy(), "temp": t[ok]}).sort_values("time")
    return {"kind": "hourly", "times": df.time.tolist(), "temps": df.temp.round(2).tolist(), "skipped_rows": int((~ok).sum())}


def _unit(m: dict, values: list[float]) -> str:
    if m.get("unit") in ("F", "C"):
        return m["unit"]
    v = np.nanmedian(values) if values else 0
    return "F" if v > 30 else "C"


def _good(out: dict) -> bool:
    n = len(out.get("dates") or out.get("times") or [])
    return n >= (3 if out["kind"] == "daily" else 72) and out["skipped_rows"] <= max(5, 0.3 * (n + out["skipped_rows"]))


# --- 4. The whole read ------------------------------------------------------------------------


def parse_file(filename: str, data: bytes) -> dict:
    rows, where = load_table(filename, data)
    rows = [r for r in rows if any((c or "").strip() for c in r)]
    if len(rows) < 3:
        raise BadRequest("That file has fewer than 3 rows of data.")
    steps = [f"Opened {where}: {len(rows):,} rows."]
    method, m, out = None, rules_mapping(rows), None
    if m:
        try:
            out = apply_mapping(rows, m)
            method = "rules" if _good(out) else None
        except BadRequest:
            out = None
    if method is None:
        if not ai_available():
            raise BadRequest("Couldn't tell which columns hold the dates and temperatures. Use columns: date, low, high "
                             "(one row per day) or date-time, temperature (a logger).")
        m = ai_mapping(rows, filename, where)
        steps.append("The layout wasn't obvious, so the AI file reader worked out which column is which.")
        out = apply_mapping(rows, m)
        if not _good(out):
            raise BadRequest("Read the file, but too few rows had a usable date and temperature. Check that it holds "
                             "daily lows and highs or timestamped readings.")
        method = "ai"
    else:
        steps.append("Recognized the columns from their names.")
    head = rows[int(m["header_rows"]) - 1] if int(m["header_rows"]) > 0 else []
    name = lambda i: (head[i] if 0 <= i < len(head) and head[i] else f"column {i + 1}") if i >= 0 else None
    values = (out.get("lows", []) + out.get("highs", [])) if out["kind"] == "daily" else out["temps"]
    unit = _unit(m, values)
    keys = out.get("dates") or out.get("times")
    first, last = keys[0][:10], keys[-1][:10]
    n = len(keys)
    if out["kind"] == "daily":
        span = (date.fromisoformat(last) - date.fromisoformat(first)).days + 1
        steps.append(f"Read {n} days of lows and highs, {first} to {last} ({span - n} days missing), in °{unit}.")
        preview = [{"when": a, "low": b, "high": c} for a, b, c in zip(out["dates"][:5], out["lows"][:5], out["highs"][:5])]
    else:
        steps.append(f"Read {n:,} logger readings, {first} to {last}, in °{unit}.")
        preview = [{"when": a, "temp": b} for a, b in zip(out["times"][:5], out["temps"][:5])]
    if out["skipped_rows"]:
        steps.append(f"Skipped {out['skipped_rows']} rows without a usable date and temperature.")
    return {**out, "unit": unit, "first": first, "last": last, "rows": n, "method": method,
            "explanation": m.get("explanation", ""), "steps": steps, "preview": preview,
            "columns": {"date": name(m["date_column"]), "time": name(m["time_column"]), "low": name(m["low_column"]),
                        "high": name(m["high_column"]), "temp": name(m["temp_column"])}}
